from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

COLUMNS = [
    ('Titre', 30, lambda s: s.title),
    ('Sous-titre', 40, lambda s: s.subtitle),
    ('Catégorie', 20, lambda s: s.categorie.name),
    ('Statut', 25, lambda s: s.get_status_display()),
    ('Lieu', 25, lambda s: s.place),
    ('Durée', 15, lambda s: s.duration.name),
    ('Âge minimum', 12, lambda s: s.minimum_age),
    ('Prix (€)', 10, lambda s: s.price),
    ("Disponible l'hiver", 12, lambda s: 'Oui' if s.available_winter else 'Non'),
    ('Pour qui ?', 40, lambda s: ', '.join(f.text for f in s.for_who.all())),
    ("Marche d'approche (min)", 15, lambda s: s.walking_time_approach),
    ('Dénivelé positif (m)', 15, lambda s: s.elevation_gain),
    ('Longueur (m)', 12, lambda s: s.distance),
    ('Description', 80, lambda s: s.description),
]


def export_sorties_xlsx(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sorties'

    for col, (header, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'A2'

    sorties = queryset.select_related('categorie', 'duration').prefetch_related('for_who')
    for row, sortie in enumerate(sorties, start=2):
        for col, (_, _, getter) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row, column=col, value=getter(sortie))
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f'sorties-{timezone.localdate():%Y-%m-%d}.xlsx'
    if len(sorties) == 1:
        filename = f'sortie-{sorties[0].slug}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
